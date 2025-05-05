import pandas as pd
from owlready2 import *
import numpy as np
import json
import glob, os
import hashlib
import datetime
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from .hlp import getXY
from datetime import datetime


def integrateReview(path):
    if "ctivit" in path:
        typeReview = "Activities"
    elif "ision" in path:
        typeReview = "Vision"
    else:
        typeReview = "ToBeConfirmed"
    wb = load_workbook(path)
    #print(wb.sheetnames)

    reviews = [x for x in wb.sheetnames if ((x not in ['Cover', 'HowTo', 'Definitions']) and (len(x) == 12))]
    if len(reviews) == 0:
        reviews = [x for x in wb.sheetnames if (x not in ['Cover', 'HowTo', 'Definitions'])]
    X, Y, terms = getXY(PATH="doc/definitions/content.md")
    inv_terms = {v: k for k, v in terms.items()}

    REVIEWS = []

    FromProbono = True
    #print(reviews)
    for R in reviews:

        ws = wb[R]
        activity = ws["C4"].value
        title = ws["C5"].value
        rev = ws["C8"].value
        origin = ws["C6"].value
        Place = ws["L6"].value
        reviewBy = ws["C7"].value
        reviewDate = ws["L7"].value

        # Parsing the table of scoring
        i = 11
        SCORES = []
        DONE = False
        #print(R)
        while not DONE:
            if not ws["C"+str(i)].value:
                DONE = True
                break
            purp = ws["C"+str(i)].value
            iss = ws["D"+str(i)].value
            purp = purp.replace("A","A.")
            iss = iss.replace("B","B.")
            scoreB = ws["E"+str(i)].value
            scoreN = ws["F"+str(i)].value
            justif = ws["G"+str(i)].value

            if scoreB:
                A = [FromProbono, origin, Place, typeReview, rev, justif, purp, iss, 'Building', scoreB, '',title, True,reviewBy,reviewDate]
                REVIEWS.append(A)
            if scoreN:
                A = [FromProbono, origin, Place, typeReview, rev, justif, purp, iss, 'Neighbourhood', scoreN, '',title, True,reviewBy,reviewDate]
                REVIEWS.append(A)
            i += 1
        print("Processing reviewed data from",path,i,typeReview)
        pdr = pd.DataFrame(REVIEWS,columns=['FromProbono', 'Origin', 'Place', 'Type', 'Source', 'Justification', 'Purpose', 'Issue', 'Scale', 'Score', 'Justification_Short', 'Source_Title', 'Reviewed',"reviewBy","reviewDate"])
        pdr["pathFile"] = path
        pdr.Purpose =pdr.Purpose.apply(lambda x: terms[x])
        pdr.Issue = pdr.Issue.apply(lambda x: terms[x])
        sign = "* "+"\n* ".join(list(pdr.Source))
        sign = str(hashlib.md5(sign.encode("utf-8")).hexdigest())
        pdr['timestamp']= datetime.now().strftime('%Y-%m-%d %X')
        pdr.to_excel("data/xls/"+sign+".xlsx")
    return pdr