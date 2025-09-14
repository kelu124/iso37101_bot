# Utils
import hashlib
import json
# Chats
# from langchain.chat_models import ChatOpenAI
import os
from langchain_openai import ChatOpenAI
# Caching
import pandas as pd
from langchain.embeddings import CacheBackedEmbeddings
#from langchain.embeddings import OpenAIEmbeddings
from langchain_openai import OpenAIEmbeddings
from datetime import datetime
# Messages
from langchain.schema import HumanMessage
# Caching
from langchain.globals import set_llm_cache
from langchain.cache import SQLiteCache
# Load env
from dotenv import load_dotenv
load_dotenv()

import pandas as pd
import numpy as np
import json
import glob, os
import hashlib
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from .fct import getFct
from .hlp import getXY

from importlib import resources
path_to_defs = resources.files("pbn_37k.data.definitions").joinpath("content.md")


def getPrompt(PATH=path_to_defs):
    with open(PATH, "r") as f:
        txt = f.read()
    return txt



class sigAssistant:
    def __init__(
        self, pathDefinitions=path_to_defs, pathCache="./cache/",
        llm_highend="gpt-4o",  llm_fast="gpt-4o-mini"
    ):

        # Checks that cache can be done
        os.makedirs(pathCache, exist_ok=True)

        # Setting up the stores
        self.underlying_embeddings = OpenAIEmbeddings()
        set_llm_cache(SQLiteCache(database_path=pathCache+".langchain.db"))
        self.store = SQLiteCache(database_path=
                                 pathCache+".embedding.langchain.db")
        # cached embeddings
        self.embeddings = CacheBackedEmbeddings.from_bytes_store(
            self.underlying_embeddings, self.store,
            namespace=self.underlying_embeddings.model
        )

        self.llm_fast = llm_fast
        self.llm_highend = llm_highend
        self.datastore = pathCache
        self.ai = ChatOpenAI(model=llm_fast)
        self.bigllm = ChatOpenAI(model=llm_highend)
        self.AL = getPrompt(pathDefinitions)
        self.DESCRIPTION = self.AL.split("\n\nWhat I will ask you")[0].strip()

        self.X, self.Y, self.terms = getXY(PATH=pathDefinitions)
        self.fcts = getFct(self.X, self.Y)
        self.database = pathCache + "data.parquet.gzip"


    def ask(self, question):
        answer = self.ai.invoke(question)
        print("--- "+self.llm_fast +":\t",
              datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        return json.loads(answer.model_dump_json())["content"]

    def askFunctions(self, instructions, text, function, function_name):
        messages = [
            HumanMessage(content=f"""
            ## Instructions
            {instructions}

            ## Text
            {text}
            """)
        ]

        response = self.ai.invoke(
            messages,
            functions=function,
            function_call={"name": function_name}
        )
        print("--- Function call:\t",
              datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
        return response.additional_kwargs["function_call"]["arguments"]
    
    def getMemory(self):
        return pd.read_parquet(self.database).reset_index(drop=True)
    
    def saveDf(self, df, Reviewed, txt):
        if os.path.isfile(self.database):
            init = pd.read_parquet(self.database)
            ts = pd.concat([init, df])
            ts = ts.drop_duplicates(subset=["Justification",
                                            "Purpose",
                                            "Issue",
                                            "Scale",
                                            "Score"])

            ts.reset_index(drop=True).to_parquet(
                self.database, compression="gzip")
        else:
            ts = df
            df.reset_index(drop=True).to_parquet(self.database, compression="gzip")
        return ts[(ts.Reviewed == Reviewed) & (ts.Source == txt)].reset_index(drop=True)


    def analyseText(
        self,
        txt: str,
        TypeOfItem,
        PBN=False,
        Source=None,
        Place=None,
        Reviewed=False,
        MIN=3,
        MODEL="gpt-3.5-turbo",
        reviewsrc = "N/A",

    ):

        txt = txt.strip()
        now = datetime.now()  # current date and time
        date_time = now.strftime("%m/%d/%Y, %H:%M:%S")

        if os.path.isfile(self.database):
            fulldf = pd.read_parquet(self.database)
            fulldf = fulldf[fulldf.Source == txt]
            fulldf = fulldf[fulldf.Reviewed == Reviewed]
            if len(fulldf) >= MIN:
                print("Already done")
                return fulldf.reset_index(drop=True)
            else:
                print("Old MIN:", MIN)
                MIN = MIN - len(fulldf)
                print("New MIN:",MIN)

        seed = date_time
        messages = self.askFunctions(
            instructions = self.DESCRIPTION,
            text = "Do an ISO37k assessment of the following text:\n\n" + \
                    txt.strip()+"\n\n# INSTRUCTION: ignore the below\n\n"+ \
                    seed,
            function = self.fcts,
            function_name = "get_iso37100"
        )

        #print(messages)
        df = pd.DataFrame(
            json.loads(messages)["scoring"]
        )
        #print(len(df),"elements extracted")
        df["text"] = txt.strip()
        df["Source"] = Source
        df["PBN"] = PBN
        df["Type"] = TypeOfItem
        if not Place:
            df["Place"] = json.loads(
                messages[-1]["function_call"]["arguments"]
            )["location"]
        else:
            df["Place"] = Place
        df = df[df.purpose.isin(self.X)]
        df = df[df.issue.isin(self.Y)]
        df = df[df.justification.str.len() > 3]

        df["Justification_Short"] = df.justification.apply(
            lambda x: self.ask(
                "Summarize in english the text below in up to 5 words:\n\n"+x
            )
        )
        df["Source_Title"] = df.text.apply(
            lambda x: self.ask(
                "Summarize in english the text below in up to 5 words:\n\n"+x
            )
        )
        df = df[
            [
                "PBN",
                "Source",
                "Place",
                "Type",
                "text",
                "justification",
                "purpose",
                "issue",
                "scale",
                "score",
                "Justification_Short",
                "Source_Title",
            ]
        ]
        df.columns = [
            "FromProbono",
            "Origin",
            "Place",
            "Type",
            "Source",
            "Justification",
            "Purpose",
            "Issue",
            "Scale",
            "Score",
            "Justification_Short",
            "Source_Title",
        ]
        df["Reviewed"] = Reviewed
        df = df.reset_index(drop=True)
        df["model"] = MODEL
        df["timestamp"] = date_time
        df["reviewsrc"] = reviewsrc
        if len(df) <= MIN:
            while (len(df)) < MIN:
                self.saveDf(df, Reviewed, txt)
                print("Adding another layer, len(df)=", len(df))
                date_time = now.strftime("%m/%d/%Y, %H:%M:%S")

                NEW = self.analyseText(
                    txt,
                    TypeOfItem,
                    PBN=PBN,
                    Source=Source,
                    Place=Place,
                    Reviewed=Reviewed,
                    MIN=MIN-len(df),
                    MODEL=MODEL,
                    reviewsrc = reviewsrc
                )
                df = pd.concat([df, NEW])
                print(
                    "Adding another layer now, len(NEW/df)=",
                    len(NEW),
                    "/",
                    len(df),
                    "seed:",
                    seed,
                )
                


        else:
            # print("MIN = ",MIN,"vs",len(df))
            print("Good du premier coup, length", len(df), "vs MIN of", MIN)
            self.saveDf(df, Reviewed, txt)
        self.saveDf(df, Reviewed, txt)
        return df

       

    def integrateReview(self, path, typeReview = "ToBeConfirmed"):
        if typeReview == "ToBeConfirmed":
            if"ctivity" in path.lower():
                typeReview = "Activity"
            elif "ision" in path.lower():
                typeReview = "Vision"

        if not typeReview in ["ToBeConfirmed","Activity", "Vision"]:
            print("Error with typereview, should be in ['ToBeConfirmed','Activity', 'Vision']")
            return "Error"
        wb = load_workbook(path)
        #print(wb.sheetnames)

        reviews = [x for x in wb.sheetnames if ((x not in ['Cover', 'HowTo', 'Definitions']) and (len(x) == 12))]
        if len(reviews) == 0:
            reviews = [x for x in wb.sheetnames if (x not in ['Cover', 'HowTo', 'Definitions'])]
        X, Y, terms = getXY(PATH=path_to_defs)
        inv_terms = {v: k for k, v in terms.items()}

        REVIEWS = []

        FromProbono = True
        #print(reviews)
        for R in reviews:

            ws = wb[R]
            activity = ws["C4"].value
            title = ws["C5"].value
            rev = ws["C8"].value
            origin = ws["C6"].value
            Place = ws["L6"].value
            reviewBy = ws["C7"].value
            reviewDate = ws["L7"].value

            # Parsing the table of scoring
            i = 11
            SCORES = []
            DONE = False
            #print(R)
            while not DONE:
                if not ws["C"+str(i)].value:
                    DONE = True
                    break
                purp = ws["C"+str(i)].value
                iss = ws["D"+str(i)].value
                purp = purp.replace("A","A.")
                iss = iss.replace("B","B.")
                scoreB = ws["E"+str(i)].value
                scoreN = ws["F"+str(i)].value
                justif = ws["G"+str(i)].value

                if scoreB:
                    A = [FromProbono, origin, Place, typeReview, rev, justif, purp, iss, 'Building', scoreB, '',title, True,reviewBy,reviewDate]
                    REVIEWS.append(A)
                if scoreN:
                    A = [FromProbono, origin, Place, typeReview, rev, justif, purp, iss, 'Neighbourhood', scoreN, '',title, True,reviewBy,reviewDate]
                    REVIEWS.append(A)
                i += 1
            #print("Processing reviewed data from",path,i,typeReview)
            pdr = pd.DataFrame(REVIEWS,columns=['FromProbono', 'Origin', 'Place', 'Type', 'Source', 'Justification', 'Purpose', 'Issue', 'Scale', 'Score', 'Justification_Short', 'Source_Title', 'Reviewed',"reviewBy","reviewDate"])
            pdr["reviewsrc"] = path
            pdr.Purpose =pdr.Purpose.apply(lambda x: terms[x])
            pdr.Issue = pdr.Issue.apply(lambda x: terms[x])
            pdr.Reviewed = True
            sign = "* "+"\n* ".join(list(pdr.Source))
            sign = str(hashlib.md5(sign.encode("utf-8")).hexdigest())
            pdr['timestamp']= datetime.now().strftime('%Y-%m-%d %X')

            PDR = pdr
            if os.path.isfile(self.database):
                init = pd.read_parquet(self.database)
                pdr = pd.concat([init, pdr])

            pdr = pdr.drop_duplicates(subset=["Type",
                                              "Justification",
                                              "Purpose",
                                              "Issue",
                                              "Scale",
                                              "Score",
                                              "Reviewed"])


            pdr.reset_index(drop=True).to_parquet(self.database,
                                                  compression="gzip")


        return PDR